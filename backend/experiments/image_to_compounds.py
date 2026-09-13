"""Small proof of concept: image -> vision model -> FooDB compounds.

The first run builds a compact SQLite index directly from the FooDB JSON archive.
Later runs only read that index.
"""

from __future__ import annotations

import argparse
import base64
import difflib
import json
import mimetypes
import os
import re
import sqlite3
import sys
import unicodedata
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

BACKEND_DIR = Path(__file__).resolve().parents[1]
DEFAULT_ARCHIVE = BACKEND_DIR / "dataset" / "raw" / "foodb" / "foodb_2020_04_07_json.zip"
DEFAULT_DB = BACKEND_DIR / "dataset" / "processed" / "foodb_compounds.sqlite3"
DEFAULT_OPENFOODTOX = (
    BACKEND_DIR / "dataset" / "raw" / "openfoodtox" / "OFT3.0_export_repository.xlsx"
)
FOODB_PREFIX = "foodb_2020_04_07_json"
MAX_IMAGE_BYTES = 15 * 1024 * 1024
COMMANDS = {"analyze", "lookup", "build-index", "build-hazards", "info"}

for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8")

console = Console()
error_console = Console(stderr=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="image_to_compounds",
        description="Image -> LiteLLM vision -> food matching -> FooDB compounds",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    query_parent = argparse.ArgumentParser(add_help=False)
    query_parent.add_argument("--db", type=Path, default=DEFAULT_DB, help="SQLite index")
    query_parent.add_argument(
        "--archive", type=Path, default=DEFAULT_ARCHIVE, help="FooDB JSON ZIP used if DB is absent"
    )
    query_parent.add_argument(
        "-n", "--max-compounds", type=int, default=20, help="Compounds shown per food"
    )
    query_parent.add_argument(
        "--min-match-score", type=float, default=0.72, help="Minimum fuzzy-match score"
    )
    query_parent.add_argument(
        "--quantified-only",
        action="store_true",
        help="Show only records with a reported content value",
    )
    query_parent.add_argument(
        "--attention",
        action="store_true",
        help="Show only quantified compounds with EFSA hazard/reference context",
    )
    query_parent.add_argument(
        "--attention-level",
        choices=("high", "watch", "all"),
        default="high",
        help="high=positive signals; watch=also ambiguous; all=also reference values",
    )
    query_parent.add_argument(
        "--openfoodtox",
        type=Path,
        default=DEFAULT_OPENFOODTOX,
        help="OpenFoodTox 3.0 Excel used by --attention",
    )
    query_parent.add_argument("--json", action="store_true", help="Print machine-readable JSON")
    query_parent.add_argument("-o", "--output", type=Path, help="Also save complete JSON here")

    analyze = subparsers.add_parser(
        "analyze",
        parents=[query_parent],
        help="Detect foods in an image and query FooDB",
        description="Send one image to the configured LiteLLM model, then query FooDB.",
    )
    analyze.add_argument("image", type=Path, help="Input JPG, PNG, WEBP, or other image")

    lookup = subparsers.add_parser(
        "lookup",
        parents=[query_parent],
        help="Query FooDB without calling the model",
        description="Look up one or more English food names directly in FooDB.",
    )
    lookup.add_argument("foods", nargs="+", help='Food names, e.g. garlic "garden onion"')

    build = subparsers.add_parser("build-index", help="Build the compact SQLite index")
    build.add_argument("--archive", type=Path, default=DEFAULT_ARCHIVE, help="FooDB JSON ZIP")
    build.add_argument("--db", type=Path, default=DEFAULT_DB, help="Output SQLite index")
    build.add_argument("--force", action="store_true", help="Replace an existing index")

    hazards = subparsers.add_parser(
        "build-hazards", help="Import a compact OpenFoodTox hazard index into SQLite"
    )
    hazards.add_argument("--source", type=Path, default=DEFAULT_OPENFOODTOX, help="OFT3 XLSX")
    hazards.add_argument("--db", type=Path, default=DEFAULT_DB, help="Target SQLite index")
    hazards.add_argument("--force", action="store_true", help="Replace existing hazard tables")

    info = subparsers.add_parser("info", help="Show local dataset/index information")
    info.add_argument("--archive", type=Path, default=DEFAULT_ARCHIVE, help="FooDB JSON ZIP")
    info.add_argument("--db", type=Path, default=DEFAULT_DB, help="SQLite index")
    info.add_argument("--openfoodtox", type=Path, default=DEFAULT_OPENFOODTOX)
    info.add_argument("--json", action="store_true", help="Print machine-readable JSON")

    argv = sys.argv[1:]
    if argv and argv[0] not in COMMANDS and not argv[0].startswith("-"):
        argv.insert(0, "analyze")
    return parser.parse_args(argv)


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def iter_json_lines(archive: zipfile.ZipFile, filename: str) -> Iterator[dict[str, Any]]:
    member = f"{FOODB_PREFIX}/{filename}"
    with archive.open(member) as source:
        for line_number, raw_line in enumerate(source, start=1):
            try:
                yield json.loads(raw_line)
            except json.JSONDecodeError as exc:
                raise RuntimeError(f"Invalid JSON in {member}:{line_number}") from exc


def create_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE food (
            id INTEGER PRIMARY KEY,
            public_id TEXT,
            name TEXT NOT NULL,
            normalized_name TEXT NOT NULL,
            scientific_name TEXT,
            food_group TEXT,
            food_subgroup TEXT
        );

        CREATE INDEX food_normalized_name_idx ON food(normalized_name);

        CREATE TABLE compound (
            id INTEGER PRIMARY KEY,
            public_id TEXT,
            name TEXT NOT NULL,
            annotation_quality TEXT,
            cas_number TEXT,
            inchikey TEXT,
            kingdom TEXT,
            superclass TEXT,
            class_name TEXT,
            subclass TEXT
        );

        CREATE TABLE food_compound (
            food_id INTEGER NOT NULL,
            compound_id INTEGER NOT NULL,
            standard_content TEXT,
            original_unit TEXT,
            food_part TEXT,
            preparation_type TEXT,
            citation TEXT,
            PRIMARY KEY (food_id, compound_id),
            FOREIGN KEY (food_id) REFERENCES food(id),
            FOREIGN KEY (compound_id) REFERENCES compound(id)
        ) WITHOUT ROWID;

        CREATE INDEX food_compound_food_idx ON food_compound(food_id);
        """
    )


def build_index(archive_path: Path, db_path: Path) -> None:
    if not archive_path.is_file():
        raise FileNotFoundError(
            f"FooDB archive not found: {archive_path}\n"
            "Download foodb_2020_04_07_json.zip from https://foodb.ca/downloads first."
        )

    db_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_db = db_path.with_suffix(f"{db_path.suffix}.building")
    temporary_db.unlink(missing_ok=True)

    error_console.print(f"[bold cyan]Building FooDB index[/]  {db_path}")
    connection = sqlite3.connect(temporary_db)
    try:
        connection.execute("PRAGMA journal_mode=OFF")
        connection.execute("PRAGMA synchronous=OFF")
        connection.execute("PRAGMA temp_store=MEMORY")
        create_schema(connection)

        with zipfile.ZipFile(archive_path) as archive:
            foods = (
                (
                    row["id"],
                    row.get("public_id"),
                    row["name"],
                    normalize_name(row["name"]),
                    row.get("name_scientific"),
                    row.get("food_group"),
                    row.get("food_subgroup"),
                )
                for row in iter_json_lines(archive, "Food.json")
                if row.get("name")
            )
            connection.executemany(
                "INSERT INTO food VALUES (?, ?, ?, ?, ?, ?, ?)",
                foods,
            )

            compounds = (
                (
                    row["id"],
                    row.get("public_id"),
                    row["name"],
                    row.get("annotation_quality"),
                    row.get("cas_number"),
                    row.get("moldb_inchikey"),
                    row.get("kingdom"),
                    row.get("superklass"),
                    row.get("klass"),
                    row.get("subklass"),
                )
                for row in iter_json_lines(archive, "Compound.json")
                if row.get("name")
            )
            connection.executemany(
                "INSERT INTO compound VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                compounds,
            )
            connection.commit()

            insert_sql = """
                INSERT OR IGNORE INTO food_compound (
                    food_id, compound_id, standard_content, original_unit,
                    food_part, preparation_type, citation
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """
            batch: list[tuple[Any, ...]] = []
            compound_rows = 0
            for row in iter_json_lines(archive, "Content.json"):
                if row.get("source_type") != "Compound":
                    continue
                food_id = row.get("food_id")
                compound_id = row.get("source_id")
                if food_id is None or compound_id is None:
                    continue
                batch.append(
                    (
                        food_id,
                        compound_id,
                        row.get("standard_content"),
                        row.get("orig_unit"),
                        row.get("orig_food_part"),
                        row.get("preparation_type"),
                        row.get("citation"),
                    )
                )
                compound_rows += 1
                if len(batch) >= 25_000:
                    connection.executemany(insert_sql, batch)
                    batch.clear()
                    if compound_rows % 500_000 == 0:
                        error_console.print(
                            f"  [dim]read {compound_rows:,} compound observations[/]"
                        )
            if batch:
                connection.executemany(insert_sql, batch)

        connection.commit()
        connection.execute("ANALYZE")
        connection.execute("VACUUM")
    except Exception:
        connection.close()
        temporary_db.unlink(missing_ok=True)
        raise
    else:
        connection.close()
        os.replace(temporary_db, db_path)

    with sqlite3.connect(db_path) as result:
        food_count = result.execute("SELECT COUNT(*) FROM food").fetchone()[0]
        compound_count = result.execute("SELECT COUNT(*) FROM compound").fetchone()[0]
        relation_count = result.execute("SELECT COUNT(*) FROM food_compound").fetchone()[0]
    error_console.print(
        "[bold green]Index ready[/]  "
        f"{food_count:,} foods · {compound_count:,} compounds · "
        f"{relation_count:,} unique relations"
    )


def sheet_records(worksheet: Any) -> Iterator[dict[str, Any]]:
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    for row in rows:
        yield {str(header): value for header, value in zip(headers, row, strict=False) if header}


def hazard_signal_level(text: str) -> str | None:
    normalized = normalize_name(text)
    high_patterns = (
        "mutagenic positive",
        "carcinogenic positive",
        "genotoxic positive",
        "carcinogenic yes",
    )
    if any(pattern in normalized for pattern in high_patterns):
        return "high"

    without_negative = normalized
    for phrase in (
        "no safety concern",
        "not a safety concern",
        "mutagenic negative",
        "carcinogenic negative",
        "genotoxic negative",
    ):
        without_negative = without_negative.replace(phrase, "")
    if any(
        pattern in without_negative
        for pattern in ("ambiguous", "equivocal", "cannot be excluded", "safety concern")
    ):
        return "watch"
    return None


def hazard_index_ready(db_path: Path) -> bool:
    if not db_path.is_file():
        return False
    with sqlite3.connect(db_path) as connection:
        exists = connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='openfoodtox_substance'"
        ).fetchone()
        if not exists:
            return False
        return connection.execute("SELECT COUNT(*) FROM openfoodtox_substance").fetchone()[0] > 0


def build_hazard_index(source_path: Path, db_path: Path, force: bool = False) -> None:
    from openpyxl import load_workbook

    if not source_path.is_file():
        raise FileNotFoundError(f"OpenFoodTox workbook not found: {source_path}")
    if not db_path.is_file():
        raise FileNotFoundError(f"FooDB SQLite index not found: {db_path}")
    if hazard_index_ready(db_path) and not force:
        error_console.print("[green]OpenFoodTox index is already ready.[/]")
        return

    error_console.print(f"[bold cyan]Importing OpenFoodTox 3.0[/]  {source_path}")
    workbook = load_workbook(source_path, read_only=True, data_only=True)

    references: dict[str, dict[str, Any]] = {}
    for row in sheet_records(workbook["REF_SUB"]):
        ref_uuid = row.get("Document UUID")
        if not ref_uuid:
            continue
        references[str(ref_uuid)] = {
            "name": row.get("ReferenceSubstanceName") or row.get("IupacName"),
            "cas_number": row.get("Inventory.CASNumber") or row.get("CAS number"),
            "inchikey": row.get("MolecularStructuralInfo.InChIKey"),
        }

    substances: dict[str, dict[str, Any]] = {}
    for row in sheet_records(workbook["SUB"]):
        sub_uuid = row.get("Document UUID")
        ref_uuid = row.get("ReferenceSubstance.ReferenceSubstance")
        if not sub_uuid:
            continue
        reference = references.get(str(ref_uuid), {})
        name = row.get("ChemicalName") or reference.get("name") or "Unknown substance"
        substances[str(sub_uuid)] = {
            "ref_uuid": str(ref_uuid) if ref_uuid else None,
            "name": name,
            "normalized_name": normalize_name(str(name)),
            "cas_number": reference.get("cas_number"),
            "inchikey": reference.get("inchikey"),
            "reference_value_count": 0,
            "endpoint_summary_count": 0,
            "high_signal_count": 0,
            "watch_signal_count": 0,
            "signals": [],
        }

    reference_markers = (
        ".RefValue.",
        ".Adi.",
        ".AcuteReferenceDose.",
        ".TolerableDailyIntake.",
        ".TolerableWeeklyIntake.",
        ".ReferenceValueDescriptor",
    )
    for row in sheet_records(workbook["FLEX_SUM.ToxRefValues"]):
        substance = substances.get(str(row.get("Parent UUID")))
        if not substance:
            continue
        has_reference = any(
            value not in (None, "") and any(marker in field for marker in reference_markers)
            for field, value in row.items()
        )
        if has_reference:
            substance["reference_value_count"] += 1

    for row in sheet_records(workbook["END_SUM"]):
        substance = substances.get(str(row.get("Parent UUID")))
        if not substance:
            continue
        key_information = str(row.get("KeyInformation.KeyInformation") or "")
        discussion = str(row.get("Discussion.Discussion") or "")
        if not key_information and not discussion:
            continue
        substance["endpoint_summary_count"] += 1
        signal_text = key_information or discussion
        level = hazard_signal_level(signal_text)
        if level == "high":
            substance["high_signal_count"] += 1
        elif level == "watch":
            substance["watch_signal_count"] += 1
        if level and len(substance["signals"]) < 5:
            substance["signals"].append(
                {
                    "level": level,
                    "endpoint": row.get("Definition"),
                    "summary": signal_text[:800],
                }
            )

    with sqlite3.connect(db_path) as connection:
        connection.executescript(
            """
            DROP TABLE IF EXISTS openfoodtox_substance;
            CREATE TABLE openfoodtox_substance (
                sub_uuid TEXT PRIMARY KEY,
                ref_uuid TEXT,
                name TEXT NOT NULL,
                normalized_name TEXT NOT NULL,
                cas_number TEXT,
                inchikey TEXT,
                reference_value_count INTEGER NOT NULL,
                endpoint_summary_count INTEGER NOT NULL,
                high_signal_count INTEGER NOT NULL,
                watch_signal_count INTEGER NOT NULL,
                signals_json TEXT NOT NULL
            );
            CREATE INDEX openfoodtox_cas_idx ON openfoodtox_substance(cas_number);
            CREATE INDEX openfoodtox_inchikey_idx ON openfoodtox_substance(inchikey);
            """
        )
        connection.executemany(
            """
            INSERT INTO openfoodtox_substance VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                (
                    sub_uuid,
                    item["ref_uuid"],
                    item["name"],
                    item["normalized_name"],
                    item["cas_number"],
                    item["inchikey"],
                    item["reference_value_count"],
                    item["endpoint_summary_count"],
                    item["high_signal_count"],
                    item["watch_signal_count"],
                    json.dumps(item["signals"], ensure_ascii=False),
                )
                for sub_uuid, item in substances.items()
            ),
        )
        connection.commit()
    workbook.close()
    matched_identifiers = sum(
        1 for item in substances.values() if item["cas_number"] or item["inchikey"]
    )
    error_console.print(
        f"[bold green]Hazard index ready[/]  {len(substances):,} substances · "
        f"{matched_identifiers:,} with CAS/InChIKey"
    )


def image_data_url(image_path: Path) -> str:
    if not image_path.is_file():
        raise FileNotFoundError(f"Image not found: {image_path}")
    size = image_path.stat().st_size
    if size > MAX_IMAGE_BYTES:
        raise ValueError(f"Image is {size / 1024 / 1024:.1f} MB; maximum is 15 MB")
    mime_type = mimetypes.guess_type(image_path.name)[0] or "image/jpeg"
    if not mime_type.startswith("image/"):
        raise ValueError(f"Unsupported image type: {mime_type}")
    encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def extract_json_object(text: str) -> dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start < 0 or end <= start:
            raise ValueError("Vision model did not return a JSON object") from None
        value = json.loads(text[start : end + 1])
    if not isinstance(value, dict):
        raise ValueError("Vision model response must be a JSON object")
    return value


def model_completion(messages: list[dict[str, Any]], temperature: float = 0.0) -> str:
    from litellm import completion

    load_dotenv(BACKEND_DIR / ".env")
    api_key = os.getenv("CUSTOM_API_KEY")
    base_url = os.getenv("CUSTOM_BASE_URL", "https://openrouter.ai/api/v1")
    model = os.getenv("CUSTOM_LLM_MODEL_NAME")
    if not api_key or not model:
        raise RuntimeError("Set CUSTOM_API_KEY and CUSTOM_LLM_MODEL_NAME in backend/.env.")

    response = completion(
        model=model,
        api_base=base_url,
        api_key=api_key,
        messages=messages,
        temperature=temperature,
    )
    content = response.choices[0].message.content
    if not isinstance(content, str):
        raise ValueError("Model returned no text response")
    return content


def detect_ingredients(image_path: Path) -> list[dict[str, Any]]:
    prompt = """
Identify only visible foods and cooking ingredients in this image. Do not infer hidden
ingredients from a finished dish. Return JSON only in this exact shape:
{"ingredients":[{"display_name":"singular common English food name",
"lookup_name":"singular common English name suitable for FooDB lookup","confidence":0.0}]}
Both name fields MUST be English, even when labels in the image use another language.
Translate visible foreign-language food labels to English. `display_name` describes what is
actually shown; `lookup_name` normalizes only true synonyms to a canonical food name.
Examples: spring onion/scallion -> green onion; cilantro -> coriander; generic edible oil ->
cooking oil. Never collapse a distinct product or plant part into its source ingredient:
bean sprout stays bean sprout (not bean), rice noodle stays rice noodle (not rice), and beef
bone stays beef bone (not beef). Use a broad name only when the image itself is uncertain.
Omit plates, tools and packaging.
""".strip()
    content = model_completion(
        [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": image_data_url(image_path)}},
                ],
            }
        ],
        temperature=0.2,
    )
    ingredients = extract_json_object(content).get("ingredients", [])
    if not isinstance(ingredients, list):
        raise ValueError("Vision model field 'ingredients' must be a list")
    return [item for item in ingredients if isinstance(item, dict) and item.get("lookup_name")]


def local_ingredients(values: list[str]) -> list[dict[str, Any]]:
    return [
        {"display_name": value.strip(), "lookup_name": value.strip(), "confidence": 1.0}
        for value in values
        if value.strip()
    ]


def best_food_match(
    connection: sqlite3.Connection, query: str, min_score: float
) -> tuple[sqlite3.Row | None, float]:
    normalized_query = normalize_name(query)
    exact = connection.execute(
        "SELECT * FROM food WHERE normalized_name = ? ORDER BY id LIMIT 1",
        (normalized_query,),
    ).fetchone()
    if exact:
        return exact, 1.0

    foods = connection.execute("SELECT * FROM food").fetchall()
    scored: list[tuple[float, sqlite3.Row, set[str]]] = []
    query_tokens = set(normalized_query.split())
    for food in foods:
        candidate = food["normalized_name"]
        score = difflib.SequenceMatcher(None, normalized_query, candidate).ratio()
        candidate_tokens = set(candidate.split())
        if query_tokens and query_tokens <= candidate_tokens:
            score = max(score, 0.9 + 0.1 * len(query_tokens) / len(candidate_tokens))
        scored.append((score, food, candidate_tokens))

    scored.sort(key=lambda item: item[0], reverse=True)
    best_score, best_row, best_tokens = scored[0]
    second_score = scored[1][0] if len(scored) > 1 else 0.0
    margin = best_score - second_score

    # A more specific detected product must not silently collapse to a broader food.
    # Examples: rice noodle -> rice, bean sprout -> bean, beef bone -> beef.
    if best_tokens < query_tokens:
        return None, best_score

    # A short term contained by many foods is ambiguous (for example generic "oil").
    if query_tokens < best_tokens:
        accepted = best_score >= max(min_score, 0.9) and margin >= 0.05
        return (best_row, best_score) if accepted else (None, best_score)

    # For non-token-equivalent names, only accept a strong, unambiguous spelling variant.
    accepted = best_score >= max(min_score, 0.84) and margin >= 0.05
    return (best_row, best_score) if accepted else (None, best_score)


def food_candidates(
    connection: sqlite3.Connection, query: str, limit: int = 20
) -> list[dict[str, Any]]:
    normalized_query = normalize_name(query)
    query_tokens = set(normalized_query.split())
    scored: list[tuple[float, sqlite3.Row]] = []
    for food in connection.execute("SELECT * FROM food").fetchall():
        candidate = food["normalized_name"]
        candidate_tokens = set(candidate.split())
        score = difflib.SequenceMatcher(None, normalized_query, candidate).ratio()
        if query_tokens and query_tokens <= candidate_tokens:
            score = max(score, 0.9 + 0.1 * len(query_tokens) / len(candidate_tokens))
        elif candidate_tokens and candidate_tokens <= query_tokens:
            score = max(score, 0.8 + 0.1 * len(candidate_tokens) / len(query_tokens))
        scored.append((score, food))
    scored.sort(key=lambda item: item[0], reverse=True)
    return [
        {
            "foodb_id": food["public_id"],
            "name": food["name"],
            "scientific_name": food["scientific_name"],
            "group": food["food_group"],
            "subgroup": food["food_subgroup"],
        }
        for _, food in scored[:limit]
    ]


def compounds_for_food(
    connection: sqlite3.Connection, food_id: int, limit: int, quantified_only: bool
) -> tuple[int, int, list[dict[str, Any]]]:
    total, quantified = connection.execute(
        """
        SELECT COUNT(*), SUM(standard_content IS NOT NULL)
        FROM food_compound
        WHERE food_id = ?
        """,
        (food_id,),
    ).fetchone()
    quantified_filter = "AND fc.standard_content IS NOT NULL" if quantified_only else ""
    rows = connection.execute(
        f"""
        SELECT c.public_id, c.name, c.annotation_quality, c.cas_number, c.inchikey,
               c.kingdom, c.superclass, c.class_name, c.subclass,
               fc.standard_content, fc.original_unit, fc.food_part,
               fc.preparation_type, fc.citation
        FROM food_compound AS fc
        JOIN compound AS c ON c.id = fc.compound_id
        WHERE fc.food_id = ?
        {quantified_filter}
        ORDER BY
            fc.standard_content IS NULL,
            CASE c.annotation_quality WHEN 'high' THEN 0 WHEN 'medium' THEN 1 ELSE 2 END,
            c.name
        LIMIT ?
        """,
        (food_id, limit),
    ).fetchall()
    return total, quantified or 0, [dict(row) for row in rows]


def attention_compounds_for_food(
    connection: sqlite3.Connection, food_id: int, limit: int, attention_level: str
) -> tuple[int, int, int, list[dict[str, Any]]]:
    total, quantified = connection.execute(
        """
        SELECT COUNT(*), SUM(
            standard_content IS NOT NULL AND CAST(standard_content AS REAL) > 0
        )
        FROM food_compound
        WHERE food_id = ?
        """,
        (food_id,),
    ).fetchone()
    rows = connection.execute(
        """
        SELECT c.public_id, c.name, c.annotation_quality, c.cas_number, c.inchikey,
               c.kingdom, c.superclass, c.class_name, c.subclass,
               fc.standard_content, fc.original_unit, fc.food_part,
               fc.preparation_type, fc.citation
        FROM food_compound AS fc
        JOIN compound AS c ON c.id = fc.compound_id
        WHERE fc.food_id = ?
          AND fc.standard_content IS NOT NULL
          AND CAST(fc.standard_content AS REAL) > 0
        """,
        (food_id,),
    ).fetchall()

    cas_values = sorted({str(row["cas_number"]) for row in rows if row["cas_number"]})
    inchikey_values = sorted({str(row["inchikey"]) for row in rows if row["inchikey"]})
    hazard_rows: list[sqlite3.Row] = []
    for field, values in (("cas_number", cas_values), ("inchikey", inchikey_values)):
        for offset in range(0, len(values), 500):
            chunk = values[offset : offset + 500]
            placeholders = ",".join("?" for _ in chunk)
            hazard_rows.extend(
                connection.execute(
                    f"SELECT * FROM openfoodtox_substance WHERE {field} IN ({placeholders})",
                    chunk,
                ).fetchall()
            )

    hazards_by_cas: dict[str, list[sqlite3.Row]] = {}
    hazards_by_inchikey: dict[str, list[sqlite3.Row]] = {}
    for hazard in hazard_rows:
        if hazard["cas_number"]:
            hazards_by_cas.setdefault(str(hazard["cas_number"]), []).append(hazard)
        if hazard["inchikey"]:
            hazards_by_inchikey.setdefault(str(hazard["inchikey"]), []).append(hazard)

    attention: list[dict[str, Any]] = []
    for row in rows:
        candidates = [
            *hazards_by_cas.get(str(row["cas_number"]), []),
            *hazards_by_inchikey.get(str(row["inchikey"]), []),
        ]
        if not candidates:
            continue
        hazard = max(
            candidates,
            key=lambda item: (
                item["high_signal_count"],
                item["watch_signal_count"],
                item["reference_value_count"],
                item["endpoint_summary_count"],
            ),
        )
        if hazard["high_signal_count"]:
            tier, score = "high_signal", 100
            why = "EFSA assessment contains a positive human-health hazard signal"
        elif hazard["watch_signal_count"] and attention_level in {"watch", "all"}:
            tier, score = "watch_signal", 70
            why = "EFSA assessment contains an ambiguous/equivocal safety signal"
        elif hazard["reference_value_count"] and attention_level == "all":
            tier, score = "reference_value", 40
            why = "EFSA has one or more toxicological reference-value records"
        else:
            continue

        compound = dict(row)
        compound["attention"] = {
            "tier": tier,
            "score": score,
            "why_selected": why,
            "authority": "EFSA OpenFoodTox 3.0",
            "openfoodtox_substance": hazard["name"],
            "reference_value_count": hazard["reference_value_count"],
            "endpoint_summary_count": hazard["endpoint_summary_count"],
            "signals": json.loads(hazard["signals_json"]),
        }
        attention.append(compound)

    attention.sort(
        key=lambda item: (
            item["attention"]["score"],
            float(item["standard_content"]),
        ),
        reverse=True,
    )
    return total, quantified or 0, len(attention), attention[:limit]


def attach_food_match(
    connection: sqlite3.Connection,
    result: dict[str, Any],
    food: sqlite3.Row,
    max_compounds: int,
    quantified_only: bool,
    attention_only: bool,
    attention_level: str,
) -> None:
    attention_count = 0
    if attention_only:
        total, quantified, attention_count, compounds = attention_compounds_for_food(
            connection, food["id"], max_compounds, attention_level
        )
    else:
        total, quantified, compounds = compounds_for_food(
            connection, food["id"], max_compounds, quantified_only
        )
    result.update(
        {
            "matched_food": {
                "foodb_id": food["public_id"],
                "name": food["name"],
                "scientific_name": food["scientific_name"],
                "group": food["food_group"],
                "subgroup": food["food_subgroup"],
            },
            "compound_count": total,
            "quantified_compound_count": quantified,
            "attention_compound_count": attention_count,
            "compounds": compounds,
        }
    )


def map_ingredients(
    db_path: Path,
    ingredients: list[dict[str, Any]],
    max_compounds: int,
    min_match_score: float,
    quantified_only: bool,
    attention_only: bool,
    attention_level: str,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    with sqlite3.connect(db_path) as connection:
        connection.row_factory = sqlite3.Row
        for ingredient in ingredients:
            query = str(ingredient["lookup_name"])
            food, score = best_food_match(connection, query, min_match_score)
            result: dict[str, Any] = {
                "detected": ingredient,
                "match_score": round(score, 3),
                "match_method": "lexical" if food else None,
                "match_relation": "exact_or_spelling" if food else None,
                "match_reason": None,
                "matched_food": None,
                "compound_count": 0,
                "quantified_compound_count": 0,
                "attention_compound_count": 0,
                "compounds": [],
            }
            if food:
                attach_food_match(
                    connection,
                    result,
                    food,
                    max_compounds,
                    quantified_only,
                    attention_only,
                    attention_level,
                )
            results.append(result)
    return results


def reconcile_unmatched_foods(
    db_path: Path,
    results: list[dict[str, Any]],
    max_compounds: int,
    quantified_only: bool,
    attention_only: bool,
    attention_level: str,
) -> None:
    unresolved: list[dict[str, Any]] = []
    allowed_candidates: dict[int, set[str]] = {}
    with sqlite3.connect(db_path) as connection:
        connection.row_factory = sqlite3.Row
        for index, result in enumerate(results):
            if result["matched_food"]:
                continue
            ingredient = result["detected"]
            candidates = food_candidates(connection, str(ingredient["lookup_name"]))
            allowed_candidates[index] = {item["foodb_id"] for item in candidates}
            unresolved.append(
                {
                    "index": index,
                    "detected_name": ingredient.get("display_name"),
                    "canonical_name": ingredient.get("lookup_name"),
                    "candidates": candidates,
                }
            )

    if not unresolved:
        return

    prompt = f"""
Resolve each detected cooking ingredient to at most one candidate from the FooDB catalog.
Return JSON only: {{"matches":[{{"index":0,"foodb_id":null,
"relation":"none","reason":"short reason"}}]}}.

Rules:
- Select only the supplied foodb_id values; otherwise use null.
- `synonym` means the same edible item under another regional/common name.
- `culinary_form` is allowed for the same botanical ingredient in an ordinary culinary form
  such as cinnamon stick -> Cinnamon, fennel seed -> Fennel, coriander seed -> Coriander,
  or garlic clove -> Garlic.
- Do NOT reduce a compositionally distinct or processed product to its source: rice noodle
  is not Rice, bean sprout is not Bean, beef bone is not Cattle, juice is not its whole fruit.
- A generic ingredient must not become an arbitrary specific variety. Generic cooking oil
  may match Cooking oil, but never Oil palm or a specific seed oil without visual evidence.
- Prefer null over a merely similar word. Relation must be `synonym`, `culinary_form`, or
  `none`.

Items and candidates:
{json.dumps(unresolved, ensure_ascii=False)}
""".strip()
    response = extract_json_object(
        model_completion([{"role": "user", "content": prompt}], temperature=0.0)
    )
    matches = response.get("matches", [])
    if not isinstance(matches, list):
        raise ValueError("Semantic matcher field 'matches' must be a list")

    with sqlite3.connect(db_path) as connection:
        connection.row_factory = sqlite3.Row
        for match in matches:
            if not isinstance(match, dict) or not isinstance(match.get("index"), int):
                continue
            index = match["index"]
            foodb_id = match.get("foodb_id")
            relation = match.get("relation")
            if (
                index not in allowed_candidates
                or not foodb_id
                or foodb_id not in allowed_candidates[index]
                or relation not in {"synonym", "culinary_form"}
            ):
                continue
            food = connection.execute(
                "SELECT * FROM food WHERE public_id = ?", (foodb_id,)
            ).fetchone()
            if not food:
                continue
            result = results[index]
            result.update(
                {
                    "match_method": "semantic",
                    "match_relation": relation,
                    "match_reason": match.get("reason"),
                }
            )
            attach_food_match(
                connection,
                result,
                food,
                max_compounds,
                quantified_only,
                attention_only,
                attention_level,
            )


def validate_query_options(args: argparse.Namespace) -> None:
    if args.max_compounds < 1:
        raise ValueError("--max-compounds must be at least 1")
    if not 0 <= args.min_match_score <= 1:
        raise ValueError("--min-match-score must be between 0 and 1")


def make_output(
    db_path: Path,
    ingredients: list[dict[str, Any]],
    max_compounds: int,
    min_match_score: float,
    quantified_only: bool,
    semantic_matching: bool,
    attention_only: bool,
    attention_level: str,
) -> dict[str, Any]:
    results = map_ingredients(
        db_path,
        ingredients,
        max_compounds,
        min_match_score,
        quantified_only,
        attention_only,
        attention_level,
    )
    if semantic_matching:
        reconcile_unmatched_foods(
            db_path,
            results,
            max_compounds,
            quantified_only,
            attention_only,
            attention_level,
        )
    return {
        "source": "FooDB 2020-04-07",
        "mode": "attention" if attention_only else "compounds",
        "note": (
            "Attention mode selects positive quantified FooDB records that can be linked by "
            "CAS/InChIKey to EFSA OpenFoodTox hazard signals or reference-value records. It is "
            "a review queue, not proof of unsafe exposure."
            if attention_only
            else "Food recognition and fuzzy matching are probabilistic. Compound records "
            "indicate reported presence in FooDB, not a complete composition or a safety "
            "conclusion."
        ),
        "results": results,
    }


def content_label(compound: dict[str, Any]) -> str:
    value = compound.get("standard_content")
    if value is None:
        return "—"
    return f"{value} {compound.get('original_unit') or ''}".strip()


def render_results(output: dict[str, Any]) -> None:
    results = output["results"]
    attention_mode = output.get("mode") == "attention"
    if not results:
        console.print(
            Panel(
                "Không phát hiện nguyên liệu hoặc thực phẩm rõ ràng trong ảnh.",
                title="[bold yellow]No food detected[/]",
                border_style="yellow",
            )
        )
        return

    summary = Table(title="Detected foods", header_style="bold cyan", show_lines=False)
    summary.add_column("Detected (English)", style="bold")
    summary.add_column("Confidence", justify="right")
    summary.add_column("FooDB match")
    summary.add_column("Match", justify="right")
    summary.add_column("Relations", justify="right")
    summary.add_column("Quantified", justify="right")
    if attention_mode:
        summary.add_column("Attention", justify="right", style="bold yellow")
    for result in results:
        detected = result["detected"]
        matched = result["matched_food"]
        confidence = detected.get("confidence")
        confidence_label = f"{float(confidence):.0%}" if confidence is not None else "—"
        display_name = str(detected.get("display_name") or detected.get("lookup_name"))
        lookup_name = str(detected.get("lookup_name") or display_name)
        detected_label = (
            display_name
            if normalize_name(display_name) == normalize_name(lookup_name)
            else f"{display_name} [dim]→ {lookup_name}[/]"
        )
        if result["match_method"] == "semantic":
            match_label = f"semantic\n[dim]{result['match_relation']}[/]"
        elif matched and result["match_score"] == 1:
            match_label = "exact"
        elif matched:
            match_label = f"lexical {result['match_score']:.0%}"
        else:
            match_label = "—"
        summary_values = [
            detected_label,
            confidence_label,
            f"{matched['name']} [dim]({matched['foodb_id']})[/]" if matched else "[red]No match[/]",
            match_label,
            f"{result['compound_count']:,}",
            f"{result['quantified_compound_count']:,}",
        ]
        if attention_mode:
            summary_values.append(f"{result['attention_compound_count']:,}")
        summary.add_row(*summary_values)
    console.print(summary)

    for result in results:
        matched = result["matched_food"]
        if not matched:
            continue
        compounds = result["compounds"]
        if attention_mode and not compounds:
            continue
        available_count = (
            result["attention_compound_count"] if attention_mode else result["compound_count"]
        )
        table = Table(
            title=(
                f"{matched['name']} · showing {len(compounds)} of "
                f"{available_count:,} "
                f"{'attention compounds' if attention_mode else 'reported compounds'}"
            ),
            header_style="bold green",
            show_lines=False,
        )
        table.add_column("FooDB ID", style="cyan", no_wrap=True)
        table.add_column("Compound", style="bold")
        table.add_column("Reported content", no_wrap=True)
        if attention_mode:
            table.add_column("EFSA attention", style="yellow")
            table.add_column("Why selected")
        else:
            table.add_column("Preparation / part")
        table.add_column("Evidence", no_wrap=True)
        for compound in compounds:
            context = " / ".join(
                value
                for value in (compound.get("preparation_type"), compound.get("food_part"))
                if value
            )
            values = [
                compound.get("public_id") or "—",
                compound["name"],
                content_label(compound),
            ]
            if attention_mode:
                attention = compound["attention"]
                values.extend([attention["tier"], attention["why_selected"]])
            else:
                values.append(context or "—")
            values.append(compound.get("citation") or "—")
            table.add_row(*values)
        console.print(table)

    console.print(f"[dim]{output['note']}[/]")


def emit_output(output: dict[str, Any], as_json: bool, output_path: Path | None) -> None:
    serialized = json.dumps(output, ensure_ascii=False, indent=2)
    if output_path:
        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(serialized + "\n", encoding="utf-8")
        error_console.print(f"[green]Saved JSON[/]  {output_path}")
    if as_json:
        print(serialized)
    else:
        render_results(output)


def database_info(archive_path: Path, db_path: Path, openfoodtox_path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "archive": {
            "path": str(archive_path.resolve()),
            "exists": archive_path.is_file(),
            "size_mb": round(archive_path.stat().st_size / 1024 / 1024, 1)
            if archive_path.is_file()
            else None,
        },
        "database": {
            "path": str(db_path.resolve()),
            "exists": db_path.is_file(),
            "size_mb": round(db_path.stat().st_size / 1024 / 1024, 1)
            if db_path.is_file()
            else None,
        },
        "openfoodtox": {
            "path": str(openfoodtox_path.resolve()),
            "exists": openfoodtox_path.is_file(),
            "size_mb": round(openfoodtox_path.stat().st_size / 1024 / 1024, 1)
            if openfoodtox_path.is_file()
            else None,
            "indexed_substances": 0,
        },
    }
    if db_path.is_file():
        with sqlite3.connect(db_path) as connection:
            result["database"].update(
                {
                    "foods": connection.execute("SELECT COUNT(*) FROM food").fetchone()[0],
                    "compounds": connection.execute("SELECT COUNT(*) FROM compound").fetchone()[0],
                    "relations": connection.execute(
                        "SELECT COUNT(*) FROM food_compound"
                    ).fetchone()[0],
                }
            )
            if connection.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='openfoodtox_substance'"
            ).fetchone():
                result["openfoodtox"]["indexed_substances"] = connection.execute(
                    "SELECT COUNT(*) FROM openfoodtox_substance"
                ).fetchone()[0]
    return result


def render_info(info: dict[str, Any]) -> None:
    table = Table(title="Local FooDB data", header_style="bold cyan")
    table.add_column("Resource", style="bold")
    table.add_column("Status")
    table.add_column("Size", justify="right")
    table.add_column("Details")
    archive = info["archive"]
    database = info["database"]
    openfoodtox = info["openfoodtox"]
    table.add_row(
        "JSON archive",
        "[green]ready[/]" if archive["exists"] else "[red]missing[/]",
        f"{archive['size_mb']} MB" if archive["size_mb"] is not None else "—",
        archive["path"],
    )
    details = "—"
    if database["exists"]:
        details = (
            f"{database['foods']:,} foods · {database['compounds']:,} compounds · "
            f"{database['relations']:,} relations"
        )
    table.add_row(
        "SQLite index",
        "[green]ready[/]" if database["exists"] else "[red]missing[/]",
        f"{database['size_mb']} MB" if database["size_mb"] is not None else "—",
        details,
    )
    hazard_status = (
        "[green]indexed[/]" if openfoodtox["indexed_substances"] else "[yellow]not indexed[/]"
    )
    table.add_row(
        "OpenFoodTox",
        hazard_status if openfoodtox["exists"] else "[red]missing[/]",
        f"{openfoodtox['size_mb']} MB" if openfoodtox["size_mb"] is not None else "—",
        f"{openfoodtox['indexed_substances']:,} indexed substances"
        if openfoodtox["indexed_substances"]
        else openfoodtox["path"],
    )
    console.print(table)


def main() -> int:
    args = parse_args()

    if args.command == "build-index":
        if args.db.is_file() and not args.force:
            error_console.print(
                f"[yellow]Index already exists:[/] {args.db.resolve()}\nUse --force to rebuild it."
            )
            return 0
        build_index(args.archive.resolve(), args.db.resolve())
        return 0

    if args.command == "build-hazards":
        build_hazard_index(args.source.resolve(), args.db.resolve(), args.force)
        return 0

    if args.command == "info":
        info = database_info(args.archive, args.db, args.openfoodtox)
        if args.json:
            print(json.dumps(info, ensure_ascii=False, indent=2))
        else:
            render_info(info)
        return 0

    validate_query_options(args)
    if not args.db.is_file():
        build_index(args.archive.resolve(), args.db.resolve())
    if args.attention and not hazard_index_ready(args.db.resolve()):
        build_hazard_index(args.openfoodtox.resolve(), args.db.resolve())

    if args.command == "analyze":
        with error_console.status(
            "[bold cyan]Detecting and resolving foods with LiteLLM...[/]", spinner="dots"
        ):
            ingredients = detect_ingredients(args.image.resolve())
            output = make_output(
                args.db.resolve(),
                ingredients,
                args.max_compounds,
                args.min_match_score,
                args.quantified_only,
                semantic_matching=True,
                attention_only=args.attention,
                attention_level=args.attention_level,
            )
    else:
        ingredients = local_ingredients(args.foods)
        output = make_output(
            args.db.resolve(),
            ingredients,
            args.max_compounds,
            args.min_match_score,
            args.quantified_only,
            semantic_matching=False,
            attention_only=args.attention,
            attention_level=args.attention_level,
        )
    emit_output(output, args.json, args.output)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, RuntimeError, ValueError, zipfile.BadZipFile) as exc:
        error_console.print(f"[bold red]Error:[/] {exc}")
        raise SystemExit(2) from exc
